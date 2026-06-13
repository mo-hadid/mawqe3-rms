#!/usr/bin/env python3
"""Export TECHNYAH_Catalog.xlsx -> products.csv + services.csv for the website.
Usage:  python export_catalog.py
Requires: pip install openpyxl
"""
import csv, sys
from openpyxl import load_workbook

SRC = "TECHNYAH_Catalog.xlsx"

def export(sheet, headers, out):
    rows = []
    started = False
    for r in sheet.iter_rows(values_only=True):
        if r and r[0] == headers[0]:        # header row
            started = True; continue
        if not started: continue
        if r is None or r[0] in (None, ""): continue
        rows.append(r[:len(headers)])
    # only published items, sorted by Sort Order
    pub_i = headers.index("Published")
    sort_i = headers.index("Sort Order")
    rows = [x for x in rows if str(x[pub_i]).strip().lower() == "yes"]
    rows.sort(key=lambda x: x[sort_i] if x[sort_i] is not None else 9999)
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for x in rows:
            w.writerow(["" if v is None else v for v in x])
    print(f"wrote {out} ({len(rows)} published items)")

try:
    wb = load_workbook(SRC, data_only=True)
except FileNotFoundError:
    sys.exit(f"Cannot find {SRC} in this folder.")

export(wb["Products"],
       ["ID","Name","Category","Status","Short Description","Image File","Sort Order","Published"],
       "products.csv")
export(wb["Services"],
       ["ID","Name","Category","Short Description","Detail (expanded)","Icon File","Sort Order","Published"],
       "services.csv")
print("Done. Upload products.csv and services.csv to the website repo.")
